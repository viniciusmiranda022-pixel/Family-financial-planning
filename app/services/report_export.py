"""Excel/PDF export for the canonical `GET /reports` payload.

This module never computes a financial fact. Every function here takes the
exact dict `app.api._build_report_payload` already returned to `GET /reports`
(the same snapshots, the same publication helpers, the same household-scoped
query) and only re-renders it into a real `.xlsx` workbook or a real `.pdf`
document -- no second parsing, classification, reconciliation, deduplication
or recalculation of any total, average, share or percentage. If a value is
wrong here, it was already wrong in `GET /reports`'s own JSON response.

Two independent renderers, one shared source of truth:

- `build_report_workbook` uses `openpyxl` (already a project dependency, used
  elsewhere for the plan-workbook importer) to write plain data cells with a
  presentation-only number format -- never a spreadsheet formula, so nothing
  here can become "the" financial engine for a value Excel might recompute
  differently than the backend did.
- `build_report_pdf` uses `pymupdf` (`Page.insert_htmlbox`, already a project
  dependency used elsewhere for PDF parsing) to render a backend-built HTML
  string through PyMuPDF's own layout engine -- no browser, no headless
  Chrome, no new dependency. This is a different mechanism from the existing
  "Imprimir / PDF" button (`window.print()` in `app/static/app.js`), which
  depends on the user's own browser; that mechanism is untouched and keeps
  working exactly as before.
"""

from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from pymupdf import Rect
from pymupdf import open as pymupdf_open

REPORT_EXPORT_MEDIA_TYPES: dict[str, str] = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}

# Friendly Portuguese labels for `report["summary"]` keys. This is presentation
# only -- a summary key missing from this map is still exported (see
# `_summary_label` below), just under its raw key name, so a future field
# added to `report_summary_monetary_publication`/`_build_report_payload`
# can never silently disappear from the export just because this map was not
# updated in lockstep.
_SUMMARY_LABELS: dict[str, str] = {
    "liquidity_name": "Nome da reserva de liquidez",
    "liquidity_direction": "Direção da liquidez no período",
    "total_spending": "Gasto total",
    "average_spending": "Gasto médio mensal",
    "total_cash_in": "Entradas totais",
    "total_cash_out": "Saídas totais",
    "total_bank_cash_out": "Saída bancária total",
    "total_card_spending": "Compras no cartão (total)",
    "cash_net": "Resultado do período",
    "savings_rate": "Taxa de poupança (%)",
    "liquidity_starting_balance": "Saldo da liquidez no início do período",
    "liquidity_balance": "Saldo da liquidez ao final do período",
    "liquidity_closing_balance": "Saldo da liquidez ao final do período (fechamento)",
    "liquidity_flow": "Fluxo líquido na liquidez no período",
    "emergency_floor": "Piso de segurança",
    "liquidity_available": "Distância do piso de segurança",
    "liquidity_deposit": "Depósito na liquidez no período",
    "liquidity_withdrawal": "Retirada da liquidez no período",
    "liquidity_uncovered_deficit": "Déficit não coberto ao final do período",
    "highest_spending": "Maior gasto mensal do período",
    "lowest_spending": "Menor gasto mensal do período",
    "highest_month": "Mês de maior gasto",
    "lowest_month": "Mês de menor gasto",
    "last_change_percentage": "Variação do último mês (%)",
}

_MONTHLY_COLUMNS: tuple[tuple[str, str], ...] = (
    ("month", "Mês"),
    ("cash_in", "Receitas reais"),
    ("bank_cash_out", "Saída bancária"),
    ("card_spending", "Compras no cartão"),
    ("spending", "Gastos"),
    ("cash_net", "Resultado"),
    ("remaining_cap", "Teto disponível"),
    ("change_percentage", "Variação (%)"),
    ("transaction_count", "Lançamentos"),
)

_CATEGORY_COLUMNS: tuple[tuple[str, str], ...] = (
    ("category", "Categoria"),
    ("amount", "Valor"),
    ("average", "Média mensal"),
    ("share", "Participação (%)"),
)

_ACCOUNT_COLUMNS: tuple[tuple[str, str], ...] = (
    ("account", "Conta"),
    ("institution", "Instituição"),
    ("account_type", "Tipo"),
    ("cash_in", "Entrou"),
    ("cash_out", "Saiu"),
    ("bank_cash_out", "Saída bancária"),
    ("card_spending", "Compras no cartão"),
    ("refunds", "Estornos"),
    ("net", "Saldo"),
)

_MONEY_FIELDS = frozenset(
    {
        "cash_in",
        "cash_out",
        "bank_cash_out",
        "card_spending",
        "spending",
        "cash_net",
        "remaining_cap",
        "cash_cap",
        "refunds",
        "net",
        "amount",
        "average",
        "total_spending",
        "average_spending",
        "total_cash_in",
        "total_cash_out",
        "total_bank_cash_out",
        "total_card_spending",
        "liquidity_starting_balance",
        "liquidity_balance",
        "liquidity_closing_balance",
        "liquidity_flow",
        "emergency_floor",
        "liquidity_available",
        "liquidity_deposit",
        "liquidity_withdrawal",
        "liquidity_uncovered_deficit",
        "highest_spending",
        "lowest_spending",
    }
)
_PERCENT_FIELDS = frozenset({"share", "change_percentage", "last_change_percentage", "savings_rate"})

_MONEY_FORMAT = '"R$" #,##0.00'
_PERCENT_FORMAT = "0.00"
_INT_FORMAT = "#,##0"


def report_export_filename(report: dict[str, Any], fmt: str) -> str:
    """A safe, non-PII filename: only the report's own period, never a
    household name, username or account identifier."""

    start = str(report.get("start_month") or "periodo")
    end = str(report.get("end_month") or "periodo")
    return f"relatorio_{start}_a_{end}.{fmt}"


def _cell_value_and_format(field: str, value: Any) -> tuple[Any, str | None]:
    if value is None:
        return "", None
    if isinstance(value, bool):
        return value, None
    if isinstance(value, (int, float, Decimal)):
        if field in _MONEY_FIELDS:
            return float(value), _MONEY_FORMAT
        if field in _PERCENT_FIELDS:
            return float(value), _PERCENT_FORMAT
        return float(value) if isinstance(value, (float, Decimal)) else int(value), _INT_FORMAT
    return str(value), None


def _write_cell(
    sheet: Worksheet,
    *,
    row: int,
    column: int,
    value: Any,
    number_format: str | None = None,
) -> Any:
    """The single place that writes a value into an openpyxl cell.

    A string value is always forced to openpyxl's plain-text data type
    (`'s'`), even when it happens to start with `=`, `+`, `-`, `@` or a
    leading tab/CR -- characters some spreadsheet clients treat as
    introducing a formula or DDE payload (openpyxl itself auto-classifies a
    leading `=` as a formula on assignment, which is exactly the P1 gap
    this closes). This never touches the *value* itself -- no evaluating,
    stripping, escaping or reinterpreting: a household-entered category,
    account or institution name of literally `=HYPERLINK(...)` is still
    shown to the household exactly as typed, just as inert text, never as
    a formula. Numeric/boolean/None values are unaffected; report data is
    never mutated (`report_export` still never computes or reclassifies a
    financial fact -- this only controls how a cell's *type* is declared).
    """

    cell = sheet.cell(row=row, column=column, value=value)
    if isinstance(value, str):
        cell.data_type = "s"
    if number_format:
        cell.number_format = number_format
    return cell


def _write_table(
    sheet: Worksheet,
    *,
    columns: tuple[tuple[str, str], ...],
    rows: list[dict[str, Any]],
) -> None:
    header_font = Font(bold=True)
    for col_index, (_field, label) in enumerate(columns, start=1):
        cell = _write_cell(sheet, row=1, column=col_index, value=label)
        cell.font = header_font
    for row_index, row in enumerate(rows, start=2):
        for col_index, (field, _label) in enumerate(columns, start=1):
            value, number_format = _cell_value_and_format(field, row.get(field))
            _write_cell(sheet, row=row_index, column=col_index, value=value, number_format=number_format)
    for col_index in range(1, len(columns) + 1):
        sheet.column_dimensions[get_column_letter(col_index)].width = 20


def build_report_workbook(report: dict[str, Any], *, generated_at: datetime) -> bytes:
    """Render `report` (the exact dict `GET /reports` returns) as a real
    `.xlsx` workbook: one sheet per canonical section, plain data cells only
    -- never a spreadsheet formula standing in for a financial calculation.
    """

    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = "Resumo"
    _write_cell(summary_sheet, row=1, column=1, value="Campo").font = Font(bold=True)
    _write_cell(summary_sheet, row=1, column=2, value="Valor").font = Font(bold=True)
    # Excel's datetime cells cannot carry tzinfo; `generated_at` is always
    # UTC (see the caller in `app/api.py`), so the label says so explicitly
    # instead of silently presenting a naive value that looks local.
    metadata_rows: list[tuple[str, Any, str | None]] = [
        ("Gerado em (UTC)", generated_at.replace(tzinfo=None), "dd/mm/yyyy hh:mm"),
        ("Período inicial", report.get("start_month"), None),
        ("Período final", report.get("end_month"), None),
        ("Meses no período", report.get("months"), _INT_FORMAT),
        ("Meses com movimentação", report.get("covered_months"), _INT_FORMAT),
        ("Duplicidades ignoradas", report.get("duplicates_ignored"), _INT_FORMAT),
    ]
    row_index = 2
    for label, value, number_format in metadata_rows:
        _write_cell(summary_sheet, row=row_index, column=1, value=label)
        _write_cell(summary_sheet, row=row_index, column=2, value=value, number_format=number_format)
        row_index += 1
    # `report["summary"]` values are not all trusted literals -- e.g.
    # `liquidity_name` echoes back a household-chosen account name, so this
    # must go through the same hardened `_write_cell` as every other
    # textual value, not a raw `.cell(...)` call.
    for field, value in report.get("summary", {}).items():
        label = _SUMMARY_LABELS.get(field, field.replace("_", " ").capitalize())
        cell_value, number_format = _cell_value_and_format(field, value)
        _write_cell(summary_sheet, row=row_index, column=1, value=label)
        _write_cell(summary_sheet, row=row_index, column=2, value=cell_value, number_format=number_format)
        row_index += 1
    summary_sheet.column_dimensions["A"].width = 34
    summary_sheet.column_dimensions["B"].width = 24

    monthly_sheet = workbook.create_sheet("Mensal")
    _write_table(monthly_sheet, columns=_MONTHLY_COLUMNS, rows=list(report.get("monthly", [])))

    categories_sheet = workbook.create_sheet("Categorias")
    _write_table(categories_sheet, columns=_CATEGORY_COLUMNS, rows=list(report.get("categories", [])))

    accounts_sheet = workbook.create_sheet("Contas")
    _write_table(accounts_sheet, columns=_ACCOUNT_COLUMNS, rows=list(report.get("accounts", [])))

    from io import BytesIO

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _escape_html(value: Any) -> str:
    text = "" if value is None else str(value)
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def _br_thousands(number: float, decimals: int) -> str:
    """Format `number` with a `.` thousands separator and `,` decimal
    separator (pt-BR convention), matching how the rest of the UI already
    presents currency (e.g. `app/static/app.js`'s own formatter). Built from
    Python's US-style `,`/`.` formatting by swapping the two separators --
    this only changes presentation, never the underlying value.
    """

    us_style = f"{number:,.{decimals}f}"  # e.g. "1,234.56"
    integer_part, _, decimal_part = us_style.partition(".")
    br_integer = integer_part.replace(",", ".")
    return f"{br_integer},{decimal_part}" if decimal_part else br_integer


def _format_cell_text(field: str, value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Sim" if value else "Não"
    if isinstance(value, (int, float, Decimal)):
        number = float(value)
        if field in _MONEY_FIELDS:
            return f"R$ {_br_thousands(number, 2)}"
        if field in _PERCENT_FIELDS:
            return f"{_br_thousands(number, 2)}%"
        if isinstance(value, int):
            return _br_thousands(number, 0)
        return _br_thousands(number, 2)
    return str(value)


def _html_table(columns: tuple[tuple[str, str], ...], rows: list[dict[str, Any]]) -> str:
    header = "".join(f"<th>{_escape_html(label)}</th>" for _field, label in columns)
    body_rows = []
    for row in rows:
        cells = "".join(
            f"<td>{_escape_html(_format_cell_text(field, row.get(field)))}</td>" for field, _label in columns
        )
        body_rows.append(f"<tr>{cells}</tr>")
    body = "".join(body_rows) or (
        f'<tr><td colspan="{len(columns)}">Sem dados no período</td></tr>'
    )
    return f"<table><thead><tr>{header}</tr></thead><tbody>{body}</tbody></table>"


_PDF_CSS = """
body { font-family: helvetica, sans-serif; font-size: 9px; color: #1f2933; }
h1 { font-size: 16px; margin: 0 0 4px 0; }
h2 { font-size: 12px; margin: 12px 0 6px 0; }
small { color: #52606d; }
table { border-collapse: collapse; width: 100%; margin-bottom: 8px; }
th, td { border: 1px solid #cbd2d9; padding: 3px 5px; text-align: left; }
th { background: #e4e7eb; }
td.right, th.right { text-align: right; }
"""

_PAGE_WIDTH = 595.0  # A4 portrait, points
_PAGE_HEIGHT = 842.0
_PAGE_MARGIN = 24.0
_MAX_ROWS_PER_PDF_PAGE = 28


def _insert_section_paginated(doc, title: str, columns: tuple[tuple[str, str], ...], rows: list[dict[str, Any]]) -> None:
    """Render one table section across as many PDF pages as it takes so no
    row is ever silently clipped -- never a single overflowing page.

    Chunked to a bounded row count per page as the primary defense (keeps
    text legible instead of relying on unbounded auto-shrink), with
    `scale_low=0` (arbitrary auto-shrink) as a second, independent defense
    and an explicit failure if a chunk still does not fit -- this function
    never returns having silently dropped content.
    """

    chunks = [rows[i : i + _MAX_ROWS_PER_PDF_PAGE] for i in range(0, len(rows), _MAX_ROWS_PER_PDF_PAGE)] or [[]]
    for chunk_index, chunk in enumerate(chunks):
        page = doc.new_page(width=_PAGE_WIDTH, height=_PAGE_HEIGHT)
        heading = title if chunk_index == 0 else f"{title} (continuação)"
        html = f"<h2>{_escape_html(heading)}</h2>" + _html_table(columns, chunk)
        rect = Rect(_PAGE_MARGIN, _PAGE_MARGIN, _PAGE_WIDTH - _PAGE_MARGIN, _PAGE_HEIGHT - _PAGE_MARGIN)
        spare_height, _scale = page.insert_htmlbox(rect, html, css=_PDF_CSS, scale_low=0)
        if spare_height == -1:
            # Should not happen with scale_low=0 (arbitrary shrink allowed),
            # but this module must never report success while silently
            # dropping rows -- fail loudly instead of returning a PDF that
            # looks complete but is missing data.
            raise RuntimeError(
                f"report_export: PDF section {heading!r} did not fit even at minimum scale "
                f"({len(chunk)} rows in this chunk)"
            )


def build_report_pdf(report: dict[str, Any], *, generated_at: datetime) -> bytes:
    """Render `report` (the exact dict `GET /reports` returns) as a real
    `.pdf`, generated server-side via PyMuPDF's own HTML/CSS layout engine
    from a backend-built HTML string -- not the user's browser, not a
    printed web page, no new dependency.
    """

    doc = pymupdf_open()

    cover_html = (
        "<h1>Relatório financeiro</h1>"
        f"<p><small>Período: {_escape_html(report.get('start_month'))} a "
        f"{_escape_html(report.get('end_month'))} &middot; "
        f"Gerado em {generated_at.strftime('%d/%m/%Y %H:%M')} UTC</small></p>"
        "<h2>Resumo</h2>"
        + _html_table(
            (("label", "Campo"), ("formatted_value", "Valor")),
            [
                {
                    "label": _SUMMARY_LABELS.get(field, field.replace("_", " ").capitalize()),
                    # Formatted using the real semantic field name (e.g.
                    # "total_spending") so money/percent fields get the right
                    # presentation, then handed to the table as an
                    # already-formatted string under a neutral column key --
                    # `_format_cell_text` would not know "total_spending"
                    # needs `R$` formatting if it only saw the generic
                    # "formatted_value" key.
                    "formatted_value": _format_cell_text(field, value),
                }
                for field, value in report.get("summary", {}).items()
            ],
        )
    )
    cover_page = doc.new_page(width=_PAGE_WIDTH, height=_PAGE_HEIGHT)
    rect = Rect(_PAGE_MARGIN, _PAGE_MARGIN, _PAGE_WIDTH - _PAGE_MARGIN, _PAGE_HEIGHT - _PAGE_MARGIN)
    spare_height, _scale = cover_page.insert_htmlbox(rect, cover_html, css=_PDF_CSS, scale_low=0)
    if spare_height == -1:
        raise RuntimeError("report_export: PDF cover/summary page did not fit even at minimum scale")

    _insert_section_paginated(doc, "Mês a mês", _MONTHLY_COLUMNS, list(report.get("monthly", [])))
    _insert_section_paginated(doc, "Categorias", _CATEGORY_COLUMNS, list(report.get("categories", [])))
    _insert_section_paginated(doc, "Contas", _ACCOUNT_COLUMNS, list(report.get("accounts", [])))

    return doc.tobytes()
