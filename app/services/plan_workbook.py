import hashlib
import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (
    Account,
    AuditEvent,
    Category,
    Commission,
    FinancialProfile,
    Household,
    Obligation,
    PayrollRecord,
    ReviewItem,
    Transaction,
    User,
)
from app.services.classifier import normalize_description
from app.services.finance import add_months, money
from app.services.importer import ParsedTransaction, transaction_fingerprint

REQUIRED_SHEETS = {
    "Plano Chácara",
    "Privilege DI",
    "Receitas e Cenários",
    "Holerites Kelly",
    "Cartão - Dados",
    "Nubank - Dados",
    "Banco - Dados",
    "Parcelas Futuras",
    "Revisar",
}

ACCOUNT_DEFINITIONS = (
    ("Itaú Conta Corrente", "Itaú", "checking", "Vinicius"),
    ("Cartão Itaú Família", "Itaú", "credit_card", "Kelly"),
    ("Cartão Nubank", "Nubank", "credit_card", "Vinicius"),
    ("Itaú Privilège DI", "Itaú", "investment", "Vinicius"),
)

CATEGORY_ALIASES = {
    "Academia": "Academia",
    "Assinaturas digitais": "Assinaturas digitais",
    "Compras e varejo": "Compras, casa e vestuário",
    "Compras, casa e vestuário": "Compras, casa e vestuário",
    "Contas do imóvel do Rodrigo": "Repasses a confirmar",
    "Créditos e estornos": "Reembolsos e estornos",
    "Dinheiro/Saque": "Pix, ajudas e outros",
    "Diversos": "Pix, ajudas e outros",
    "Entrada da chácara (provável)": "Operação da chácara",
    "Estética e beleza": "Estética e beleza",
    "Estorno recebido": "Reembolsos e estornos",
    "Juros, IOF e tarifas": "Juros, IOF e tarifas",
    "Mercado e itens domésticos": "Mercado e itens domésticos",
    "Movimentação de investimento": "Transferência patrimonial",
    "Outras entradas/transferências": "Repasses a confirmar",
    "Outras faturas/crediário": "Pix, ajudas e outros",
    "Outros débitos": "Pix, ajudas e outros",
    "Pagamento fatura Itaú": "Conciliação",
    "Pagamento fatura Nubank": "Conciliação",
    "Pagamento recebido": "Conciliação",
    "Restaurantes e delivery": "Restaurantes e delivery",
    "Saúde": "Saúde e farmácia",
    "Saúde e farmácia": "Saúde e farmácia",
    "Saúde/Casa a revisar": "Revisar",
    "Seguros": "Seguros",
    "Telefonia e internet": "Água e energia da residência",
    "Transferência entre contas próprias": "Transferência interna",
    "Transferência familiar - Kelly": "Transferência interna",
    "Transferências a pessoas": "Pix, ajudas e outros",
    "Transferências grandes a identificar": "Repasses a confirmar",
    "Transporte": "Transporte",
    "Viagens e lazer": "Viagens e lazer",
    "Água e energia da residência": "Água e energia da residência",
}

ENVELOPE_CATEGORY_MAP = {
    "Mercado + refeições": (
        ("Mercado e itens domésticos", False),
        ("Restaurantes e delivery", False),
    ),
    "Casa: CPFL, BRK, internet, telefone e seguros": (
        ("Água e energia da residência", True),
    ),
    "Transporte": (("Transporte", True),),
    "Saúde e farmácia": (("Saúde e farmácia", True),),
    "Academia": (("Academia", False),),
    "Assinaturas": (("Assinaturas digitais", False),),
    "Compras, lazer e beleza": (
        ("Compras, casa e vestuário", False),
        ("Viagens e lazer", False),
        ("Estética e beleza", False),
    ),
    "Pix, ajudas e outros": (("Pix, ajudas e outros", False),),
    "Operação da chácara": (("Operação da chácara", True),),
    "Margem para imprevistos": (("Margem para imprevistos", True),),
}

CATEGORY_COLORS = {
    "Conciliação": "#64748B",
    "Transferência patrimonial": "#64748B",
    "Transferência interna": "#64748B",
    "Repasses a confirmar": "#F59E0B",
    "Reembolsos e estornos": "#22C55E",
    "Restaurantes e delivery": "#F97316",
    "Mercado e itens domésticos": "#10B981",
    "Água e energia da residência": "#0EA5E9",
    "Transporte": "#3B82F6",
    "Saúde e farmácia": "#EF4444",
    "Academia": "#8B5CF6",
    "Assinaturas digitais": "#6366F1",
    "Estética e beleza": "#EC4899",
    "Viagens e lazer": "#F43F5E",
    "Seguros": "#14B8A6",
    "Juros, IOF e tarifas": "#DC2626",
    "Compras, casa e vestuário": "#A855F7",
    "Pix, ajudas e outros": "#D97706",
    "Operação da chácara": "#15803D",
    "Margem para imprevistos": "#475569",
    "Revisar": "#F59E0B",
}


@dataclass(frozen=True)
class ProfileData:
    monthly_salary_net: Decimal
    monthly_cash_cap: Decimal
    emergency_floor: Decimal
    food_allowance: Decimal
    meal_allowance_daily: Decimal
    workdays_month: int
    investment_name: str
    investment_balance: Decimal
    investment_gross_annual_rate: Decimal
    investment_income_tax_rate: Decimal
    projection_end: date


@dataclass(frozen=True)
class PayrollData:
    person_name: str
    competence: date
    payment_date: date
    payroll_kind: str
    gross_amount: Decimal
    deductions: Decimal
    net_amount: Decimal
    payroll_loan: Decimal
    notes: str


@dataclass(frozen=True)
class CommissionData:
    description: str
    expected_date: date
    gross_amount: Decimal
    tax_rate: Decimal
    delay_days: int


@dataclass(frozen=True)
class ObligationData:
    name: str
    due_date: date
    amount: Decimal
    recurrence_months: int
    occurrence_count: int
    category: str


@dataclass(frozen=True)
class WorkbookTransaction:
    account_name: str
    booked_at: date
    description: str
    amount: Decimal
    category: str
    transaction_type: str
    owner_label: str
    card_last_four: str | None
    installment_current: int | None
    installment_total: int | None
    source_line: int
    excluded: bool
    confidence: Decimal
    review_reason: str | None


@dataclass(frozen=True)
class PlanReviewData:
    reason: str
    details: str


@dataclass(frozen=True)
class PlanWorkbookData:
    profile: ProfileData
    category_caps: dict[str, tuple[Decimal | None, bool]]
    payroll: tuple[PayrollData, ...]
    commissions: tuple[CommissionData, ...]
    obligations: tuple[ObligationData, ...]
    transactions: tuple[WorkbookTransaction, ...]
    reviews: tuple[PlanReviewData, ...]


@dataclass
class ImportStats:
    created: dict[str, int] = field(default_factory=lambda: defaultdict(int))
    updated: dict[str, int] = field(default_factory=lambda: defaultdict(int))
    unchanged: dict[str, int] = field(default_factory=lambda: defaultdict(int))

    def mark(self, entity: str, state: str) -> None:
        getattr(self, state)[entity] += 1

    def as_dict(self) -> dict[str, dict[str, int]]:
        entities = sorted({*self.created, *self.updated, *self.unchanged})
        return {
            entity: {
                "created": self.created[entity],
                "updated": self.updated[entity],
                "unchanged": self.unchanged[entity],
            }
            for entity in entities
        }


def _decimal(value: object) -> Decimal:
    if value in (None, ""):
        return Decimal("0.00")
    return money(value)


def _rate(value: object) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    return Decimal(str(value)).quantize(Decimal("0.00000001"))


def _date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        normalized = value.strip()
        if re.fullmatch(r"\d{4}-\d{2}", normalized):
            normalized = f"{normalized}-01"
        return datetime.fromisoformat(normalized[:10]).date()
    raise ValueError(f"Data inválida na planilha: {value!r}")


def _int_or_none(value: object) -> int | None:
    return int(value) if value not in (None, "") else None


def _bool(value: object) -> bool:
    if isinstance(value, str):
        return value.strip().lower() in {"true", "sim", "1", "yes"}
    return bool(value)


def _label_values(sheet, start: int, end: int, label_col: int = 1, value_col: int = 2) -> dict[str, object]:
    values: dict[str, object] = {}
    for row in range(start, end + 1):
        label = sheet.cell(row, label_col).value
        if label:
            values[str(label).strip()] = sheet.cell(row, value_col).value
    return values


def _category_name(value: object) -> str:
    raw = str(value or "Revisar").strip()
    return CATEGORY_ALIASES.get(raw, raw)


def _classification(category: str, nature: str, amount: Decimal, count_in_spend: bool = True) -> tuple[str, bool, Decimal, str | None]:
    normalized_nature = normalize_description(nature)
    if category == "Conciliação":
        return "reconciliation", True, Decimal("0.99"), None
    if category in {"Transferência patrimonial", "Transferência interna"}:
        return "transfer", True, Decimal("0.97"), None
    if "EXCLUIR" in normalized_nature or not count_in_spend:
        transaction_type = "income" if amount > 0 else "expense"
        return transaction_type, True, Decimal("0.90"), None
    if category == "Reembolsos e estornos" or amount > 0 and "ESTORNO" in normalized_nature:
        return "refund", False, Decimal("0.98"), None
    transaction_type = "income" if amount > 0 else "expense"
    return transaction_type, False, Decimal("0.95"), None


def _card_amount(value: object) -> Decimal:
    raw = _decimal(value)
    return abs(raw) if raw < 0 else -abs(raw)


def _card_owner(holder: object, default: str) -> tuple[str, str | None]:
    text = str(holder or "").strip()
    last_four = None
    match = re.search(r"final\s+(\d{4})", text, re.IGNORECASE)
    if match:
        last_four = match.group(1)
    normalized = normalize_description(text)
    if "VINICIUS" in normalized:
        return "Vinicius", last_four
    if "KELLY" in normalized:
        return "Kelly", last_four
    return default, last_four


def _installment_base(description: str) -> str:
    without_installment = re.sub(
        r"(?:PARCELA\s*)?\d{1,2}\s*/\s*\d{1,2}", " ", description, flags=re.IGNORECASE
    )
    return normalize_description(without_installment)


def _nubank_installment_anchors(workbook) -> dict[tuple[str, int], date]:
    anchors: dict[tuple[str, int], date] = {}
    sheet = workbook["Parcelas Futuras"]
    for row in sheet.iter_rows(min_row=6, values_only=True):
        due_month, issuer, merchant, total = row[5], row[6], row[7], row[9]
        if not due_month or str(issuer or "").strip().lower() != "nubank" or not total:
            continue
        key = (_installment_base(str(merchant)), int(total))
        due = _date(due_month).replace(day=1)
        anchors[key] = min(anchors.get(key, due), due)
    return {key: add_months(due, -1) for key, due in anchors.items()}


def _parse_review_sheet(workbook) -> tuple[tuple[PlanReviewData, ...], set[tuple[date, str, Decimal]]]:
    sheet = workbook["Revisar"]
    reviews: list[PlanReviewData] = []
    for row in sheet.iter_rows(min_row=6, max_row=22, values_only=True):
        priority, question, impact, status = row[:4]
        if str(status or "").strip().lower() != "aguardando" or not question:
            continue
        digest = hashlib.sha1(str(question).encode()).hexdigest()[:12]
        reviews.append(
            PlanReviewData(
                reason=f"plan_question:{digest}",
                details=f"Prioridade {priority}: {question} Impacto: {impact}",
            )
        )
    transaction_keys: set[tuple[date, str, Decimal]] = set()
    for row in sheet.iter_rows(min_row=27, max_row=80, values_only=True):
        if not row or not row[0] or not row[1]:
            continue
        transaction_keys.add((_date(row[0]), normalize_description(str(row[1])), _decimal(row[2])))
    return tuple(reviews), transaction_keys


def _parse_transactions(
    workbook, transaction_review_keys: set[tuple[date, str, Decimal]]
) -> tuple[WorkbookTransaction, ...]:
    result: list[WorkbookTransaction] = []

    card_sheet = workbook["Cartão - Dados"]
    for row_number, row in enumerate(card_sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row[0] or not row[5]:
            continue
        amount = _card_amount(row[6])
        category = _category_name(row[7])
        nature = str(row[8] or "")
        transaction_type, excluded, confidence, review = _classification(category, nature, amount)
        if category == "Revisar":
            confidence = Decimal("0.55")
            review = "Classificação indicada para revisão na planilha"
        owner, last_four = _card_owner(row[3], "Kelly")
        result.append(
            WorkbookTransaction(
                account_name="Cartão Itaú Família",
                booked_at=_date(row[0]),
                description=str(row[5]).strip(),
                amount=amount,
                category=category,
                transaction_type=transaction_type,
                owner_label=owner,
                card_last_four=last_four,
                installment_current=_int_or_none(row[10]),
                installment_total=_int_or_none(row[11]),
                source_line=row_number,
                excluded=excluded,
                confidence=confidence,
                review_reason=review,
            )
        )

    nubank_anchors = _nubank_installment_anchors(workbook)
    nubank_sheet = workbook["Nubank - Dados"]
    for row_number, row in enumerate(nubank_sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row[0] or not row[2]:
            continue
        description = str(row[2]).strip()
        current, total = _int_or_none(row[7]), _int_or_none(row[8])
        booked_at = _date(row[0])
        if current and total:
            booked_at = nubank_anchors.get((_installment_base(description), total), booked_at)
        amount = _card_amount(row[3])
        category = _category_name(row[4])
        nature = str(row[5] or "")
        transaction_type, excluded, confidence, review = _classification(category, nature, amount)
        result.append(
            WorkbookTransaction(
                account_name="Cartão Nubank",
                booked_at=booked_at,
                description=description,
                amount=amount,
                category=category,
                transaction_type=transaction_type,
                owner_label="Vinicius",
                card_last_four=None,
                installment_current=current,
                installment_total=total,
                source_line=row_number,
                excluded=excluded,
                confidence=confidence,
                review_reason=review,
            )
        )

    bank_sheet = workbook["Banco - Dados"]
    for row_number, row in enumerate(bank_sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row[0] or not row[2]:
            continue
        amount = _decimal(row[3])
        category = _category_name(row[4])
        nature = str(row[5] or "")
        transaction_type, excluded, confidence, review = _classification(
            category, nature, amount, _bool(row[8])
        )
        transaction_key = (_date(row[0]), normalize_description(str(row[2])), amount)
        if transaction_key in transaction_review_keys:
            confidence = Decimal("0.55")
            review = "Transferência extraordinária destacada para confirmação na planilha"
        result.append(
            WorkbookTransaction(
                account_name="Itaú Conta Corrente",
                booked_at=_date(row[0]),
                description=str(row[2]).strip(),
                amount=amount,
                category=category,
                transaction_type=transaction_type,
                owner_label="Vinicius",
                card_last_four=None,
                installment_current=None,
                installment_total=None,
                source_line=row_number,
                excluded=excluded,
                confidence=confidence,
                review_reason=review,
            )
        )
    return tuple(result)


def parse_plan_workbook(path: str | Path) -> PlanWorkbookData:
    source = Path(path)
    if not source.is_file():
        raise ValueError(f"Planilha não encontrada: {source}")
    workbook = load_workbook(source, read_only=True, data_only=True)
    missing = REQUIRED_SHEETS.difference(workbook.sheetnames)
    if missing:
        raise ValueError(f"Planilha incompatível; abas ausentes: {', '.join(sorted(missing))}")

    plan = workbook["Plano Chácara"]
    assumptions = _label_values(plan, 7, 24)
    privilege = _label_values(workbook["Privilege DI"], 7, 15)

    projection_months = [
        _date(workbook["Receitas e Cenários"].cell(row, 1).value)
        for row in range(16, 80)
        if workbook["Receitas e Cenários"].cell(row, 1).value
    ]
    if not projection_months:
        raise ValueError("A projeção mensal não foi encontrada na planilha")
    profile = ProfileData(
        monthly_salary_net=_decimal(assumptions["Kelly - salário líquido base"]),
        monthly_cash_cap=_decimal(assumptions["Teto mensal recomendado em dinheiro"]),
        emergency_floor=_decimal(assumptions["Reserva de emergência alvo"]),
        food_allowance=_decimal(assumptions["Vale-alimentação mensal"]),
        meal_allowance_daily=_decimal(assumptions["Vale-refeição por dia"]),
        workdays_month=int(assumptions["Dias trabalhados/mês"]),
        investment_name="Itaú Privilège RF Referenciado DI",
        investment_balance=_decimal(assumptions["Dinheiro investido hoje"]),
        investment_gross_annual_rate=_rate(privilege["Retorno bruto anual estimado"]),
        investment_income_tax_rate=_rate(privilege["IR conservador sobre os ganhos"]),
        projection_end=max(projection_months),
    )

    category_caps: dict[str, tuple[Decimal | None, bool]] = {}
    for row in plan.iter_rows(min_row=28, max_row=37, values_only=True):
        envelope = str(row[0] or "").strip()
        if envelope not in ENVELOPE_CATEGORY_MAP:
            continue
        cap = _decimal(row[2])
        mapped = ENVELOPE_CATEGORY_MAP[envelope]
        primary = True
        for category_name, essential in mapped:
            category_caps[category_name] = (cap if primary else Decimal("0.00"), essential)
            primary = False

    payroll: list[PayrollData] = []
    payroll_sheet = workbook["Holerites Kelly"]
    for row in payroll_sheet.iter_rows(min_row=6, max_row=8, values_only=True):
        payroll.append(
            PayrollData(
                person_name="Kelly Cristina Rodrigues Miranda",
                competence=_date(row[0]).replace(day=1),
                payment_date=_date(row[1]),
                payroll_kind="regular",
                gross_amount=_decimal(row[2]),
                deductions=_decimal(row[3]),
                net_amount=_decimal(row[4]),
                payroll_loan=_decimal(row[6]),
                notes=f"Valor real consolidado na planilha; fonte: {row[7]}",
            )
        )
    payroll.extend(
        (
            PayrollData(
                "Kelly Cristina Rodrigues Miranda",
                date(2026, 11, 1),
                date(2026, 11, 30),
                "13_first",
                _decimal(payroll_sheet["B27"].value),
                Decimal("0.00"),
                _decimal(payroll_sheet["B27"].value),
                Decimal("0.00"),
                "Estimativa da 1ª parcela do 13º; substituir quando o demonstrativo for emitido.",
            ),
            PayrollData(
                "Kelly Cristina Rodrigues Miranda",
                date(2026, 12, 1),
                date(2026, 12, 20),
                "13_second",
                _decimal(payroll_sheet["B28"].value),
                Decimal("0.00"),
                _decimal(payroll_sheet["B28"].value),
                Decimal("0.00"),
                "Estimativa líquida da 2ª parcela do 13º; substituir pelo valor real.",
            ),
            PayrollData(
                "Kelly Cristina Rodrigues Miranda",
                date(2026, 12, 1),
                date(2026, 12, 1),
                "vacation_extra",
                _decimal(payroll_sheet["G21"].value),
                _decimal(payroll_sheet["G23"].value),
                _decimal(payroll_sheet["G24"].value),
                Decimal("0.00"),
                "Estimativa: somente o ganho adicional líquido do terço de 15 dias; sem duplicar salário.",
            ),
        )
    )

    commissions: list[CommissionData] = []
    income_sheet = workbook["Receitas e Cenários"]
    tax_rate = _rate(income_sheet["K6"].value)
    for row in income_sheet.iter_rows(min_row=7, max_row=10, values_only=True):
        commissions.append(
            CommissionData(
                description=f"{row[0]} — Recebível PJ",
                expected_date=_date(row[1]).replace(day=1),
                gross_amount=_decimal(row[3]),
                tax_rate=tax_rate,
                delay_days=60,
            )
        )

    flow_rows = []
    for row in plan.iter_rows(min_row=42, max_row=100, values_only=True):
        if not row or not row[0]:
            break
        flow_rows.append(row)
    installment_rows = [row for row in flow_rows if _decimal(row[5]) > 0]
    obligations: list[ObligationData] = []
    if installment_rows:
        obligations.append(
            ObligationData(
                name="Chácara — parcelas mensais",
                due_date=_date(installment_rows[0][0]).replace(day=10),
                amount=_decimal(installment_rows[0][5]),
                recurrence_months=1,
                occurrence_count=len(installment_rows),
                category="chacara",
            )
        )
    reinforcement_number = 0
    for row in flow_rows:
        amount = _decimal(row[6])
        if amount <= 0:
            continue
        reinforcement_number += 1
        due = _date(row[0]).replace(day=10)
        obligations.append(
            ObligationData(
                name=f"Chácara — reforço {reinforcement_number} ({due:%m/%Y})",
                due_date=due,
                amount=amount,
                recurrence_months=0,
                occurrence_count=1,
                category="chacara",
            )
        )

    reviews, transaction_review_keys = _parse_review_sheet(workbook)
    transactions = _parse_transactions(workbook, transaction_review_keys)
    workbook.close()
    return PlanWorkbookData(
        profile=profile,
        category_caps=category_caps,
        payroll=tuple(payroll),
        commissions=tuple(commissions),
        obligations=tuple(obligations),
        transactions=transactions,
        reviews=reviews,
    )


def select_household(db: Session, household_name: str | None = None) -> Household:
    query = select(Household)
    if household_name:
        query = query.where(Household.name == household_name)
    households = db.scalars(query.order_by(Household.created_at)).all()
    if not households:
        raise ValueError("Nenhuma família cadastrada. Conclua o primeiro acesso pelo navegador.")
    if len(households) > 1:
        raise ValueError("Há mais de uma família; informe --household com o nome exato.")
    return households[0]


def _set_fields(target: object, values: dict[str, object]) -> bool:
    changed = False
    for field_name, value in values.items():
        if getattr(target, field_name) != value:
            setattr(target, field_name, value)
            changed = True
    return changed


def _mark_upsert(stats: ImportStats, entity: str, created: bool, changed: bool = False) -> None:
    stats.mark(entity, "created" if created else "updated" if changed else "unchanged")


def import_plan_data(
    db: Session,
    household: Household,
    data: PlanWorkbookData,
    document_id: str | None = None,
) -> ImportStats:
    stats = ImportStats()
    household_id = household.id

    profile = db.scalar(select(FinancialProfile).where(FinancialProfile.household_id == household_id))
    profile_values = data.profile.__dict__
    if profile is None:
        profile = FinancialProfile(household_id=household_id, **profile_values)
        db.add(profile)
        _mark_upsert(stats, "profile", True)
    else:
        _mark_upsert(stats, "profile", False, _set_fields(profile, profile_values))

    account_by_name = {
        item.name: item
        for item in db.scalars(select(Account).where(Account.household_id == household_id)).all()
    }
    for name, institution, account_type, owner_label in ACCOUNT_DEFINITIONS:
        values = {
            "institution": institution,
            "account_type": account_type,
            "owner_label": owner_label,
            "active": True,
        }
        account = account_by_name.get(name)
        if account is None:
            account = Account(household_id=household_id, name=name, **values)
            db.add(account)
            db.flush()
            account_by_name[name] = account
            _mark_upsert(stats, "accounts", True)
        else:
            _mark_upsert(stats, "accounts", False, _set_fields(account, values))

    category_by_name = {
        item.name: item
        for item in db.scalars(select(Category).where(Category.household_id == household_id)).all()
    }
    category_names = {_category_name(item.category) for item in data.transactions}
    category_names.update(data.category_caps)
    for name in sorted(category_names):
        cap, essential = data.category_caps.get(name, (None, False))
        values = {
            "color": CATEGORY_COLORS.get(name, "#64748B"),
            "cash_cap": cap,
            "essential": essential,
        }
        category = category_by_name.get(name)
        if category is None:
            category = Category(household_id=household_id, name=name, **values)
            db.add(category)
            db.flush()
            category_by_name[name] = category
            _mark_upsert(stats, "categories", True)
        else:
            _mark_upsert(stats, "categories", False, _set_fields(category, values))

    existing_commissions = {
        item.description: item
        for item in db.scalars(select(Commission).where(Commission.household_id == household_id)).all()
    }
    for source in data.commissions:
        values = {
            "expected_date": source.expected_date,
            "gross_amount": source.gross_amount,
            "tax_rate": source.tax_rate,
            "delay_days": source.delay_days,
            "status": "expected",
        }
        item = existing_commissions.get(source.description)
        if item is None:
            item = Commission(household_id=household_id, description=source.description, **values)
            db.add(item)
            _mark_upsert(stats, "commissions", True)
        else:
            _mark_upsert(stats, "commissions", False, _set_fields(item, values))

    payroll_by_key = {
        (item.person_name, item.competence, item.payroll_kind): item
        for item in db.scalars(select(PayrollRecord).where(PayrollRecord.household_id == household_id)).all()
    }
    for source in data.payroll:
        key = (source.person_name, source.competence, source.payroll_kind)
        values = {
            "document_id": document_id,
            "payment_date": source.payment_date,
            "gross_amount": source.gross_amount,
            "deductions": source.deductions,
            "net_amount": source.net_amount,
            "payroll_loan": source.payroll_loan,
            "notes": source.notes,
        }
        item = payroll_by_key.get(key)
        if item is None:
            item = PayrollRecord(
                household_id=household_id,
                person_name=source.person_name,
                competence=source.competence,
                payroll_kind=source.payroll_kind,
                **values,
            )
            db.add(item)
            _mark_upsert(stats, "payroll", True)
        else:
            _mark_upsert(stats, "payroll", False, _set_fields(item, values))

    obligation_by_name = {
        item.name: item
        for item in db.scalars(select(Obligation).where(Obligation.household_id == household_id)).all()
    }
    for source in data.obligations:
        values = {
            "due_date": source.due_date,
            "amount": source.amount,
            "recurrence_months": source.recurrence_months,
            "occurrence_count": source.occurrence_count,
            "category": source.category,
            "active": True,
        }
        item = obligation_by_name.get(source.name)
        if item is None:
            item = Obligation(household_id=household_id, name=source.name, **values)
            db.add(item)
            _mark_upsert(stats, "obligations", True)
        else:
            _mark_upsert(stats, "obligations", False, _set_fields(item, values))

    existing_transactions = {
        item.fingerprint: item
        for item in db.scalars(select(Transaction).where(Transaction.household_id == household_id)).all()
    }
    occurrence_counts: dict[str, int] = defaultdict(int)
    for source in data.transactions:
        account = account_by_name[source.account_name]
        parsed = ParsedTransaction(
            booked_at=source.booked_at,
            description=source.description,
            amount=source.amount,
            source_line=source.source_line,
            card_last_four=source.card_last_four,
            installment_current=source.installment_current,
            installment_total=source.installment_total,
        )
        canonical = transaction_fingerprint(account.id, parsed, source.owner_label)
        occurrence = occurrence_counts[canonical]
        occurrence_counts[canonical] += 1
        fingerprint = canonical
        if occurrence:
            fingerprint = hashlib.sha256(f"{canonical}|plan-occurrence:{occurrence}".encode()).hexdigest()
        values = {
            "account_id": account.id,
            "document_id": document_id,
            "category_id": category_by_name[source.category].id,
            "booked_at": source.booked_at,
            "description": source.description,
            "normalized_description": normalize_description(source.description),
            "amount": source.amount,
            "transaction_type": source.transaction_type,
            "owner_label": source.owner_label,
            "card_last_four": source.card_last_four,
            "installment_current": source.installment_current,
            "installment_total": source.installment_total,
            "source_line": source.source_line,
            "confidence": source.confidence,
            "excluded": source.excluded,
            "possible_duplicate": False,
        }
        item = existing_transactions.get(fingerprint)
        if item is None:
            item = Transaction(household_id=household_id, fingerprint=fingerprint, **values)
            db.add(item)
            db.flush()
            existing_transactions[fingerprint] = item
            _mark_upsert(stats, "transactions", True)
        else:
            _mark_upsert(stats, "transactions", False, _set_fields(item, values))
        if source.review_reason:
            review = db.scalar(
                select(ReviewItem).where(
                    ReviewItem.transaction_id == item.id,
                    ReviewItem.reason == "spreadsheet_review",
                )
            )
            if review is None:
                db.add(
                    ReviewItem(
                        household_id=household_id,
                        transaction_id=item.id,
                        document_id=document_id,
                        reason="spreadsheet_review",
                        details=source.review_reason,
                    )
                )
                _mark_upsert(stats, "reviews", True)
            else:
                changed = _set_fields(
                    review,
                    {"document_id": document_id, "details": source.review_reason},
                )
                _mark_upsert(stats, "reviews", False, changed)

    existing_questions = {
        item.reason: item
        for item in db.scalars(
            select(ReviewItem).where(
                ReviewItem.household_id == household_id,
                ReviewItem.reason.like("plan_question:%"),
            )
        ).all()
    }
    for source in data.reviews:
        review = existing_questions.get(source.reason)
        values = {"document_id": document_id, "details": source.details}
        if review is None:
            db.add(
                ReviewItem(
                    household_id=household_id,
                    reason=source.reason,
                    **values,
                )
            )
            _mark_upsert(stats, "reviews", True)
        else:
            _mark_upsert(stats, "reviews", False, _set_fields(review, values))

    db.flush()
    return stats


def add_import_audit(db: Session, household: Household, stats: ImportStats, filename: str) -> None:
    user = db.scalar(
        select(User)
        .where(User.household_id == household.id, User.active.is_(True))
        .order_by(User.is_admin.desc(), User.created_at)
    )
    db.add(
        AuditEvent(
            household_id=household.id,
            user_id=user.id if user else None,
            event_type="financial_plan.import",
            entity_type="financial_plan_workbook",
            details=str({"filename": filename, "stats": stats.as_dict()}),
        )
    )
