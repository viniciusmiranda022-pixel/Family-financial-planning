import os
import uuid
from datetime import date
from decimal import Decimal

from cryptography.fernet import Fernet
from openpyxl import Workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session

os.environ.setdefault("DATABASE_URL", f"sqlite:////tmp/ffp-plan-{uuid.uuid4().hex}.sqlite")
os.environ.setdefault("DATA_DIR", f"/tmp/ffp-plan-data-{uuid.uuid4().hex}")
os.environ.setdefault("SECRET_KEY", "plan-import-test-secret-that-is-long-enough")
os.environ.setdefault("FILE_ENCRYPTION_KEY", Fernet.generate_key().decode())

from app.api import _future_installments  # noqa: E402
from app.db import Base  # noqa: E402
from app.models import Household, Transaction  # noqa: E402
from app.services.plan_workbook import import_plan_data, parse_plan_workbook  # noqa: E402


def _minimal_plan(path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in (
        "Plano Chácara",
        "Privilege DI",
        "Receitas e Cenários",
        "Holerites Kelly",
        "Cartão - Dados",
        "Nubank - Dados",
        "Banco - Dados",
        "Parcelas Futuras",
        "Revisar",
    ):
        workbook.create_sheet(name)

    plan = workbook["Plano Chácara"]
    assumptions = (
        (7, "Dinheiro investido hoje", 43875.62),
        (8, "Kelly - salário líquido base", 5486.36),
        (11, "Imposto sobre as comissões", 0.06),
        (13, "Vale-alimentação mensal", 1950),
        (14, "Vale-refeição por dia", 43),
        (15, "Dias trabalhados/mês", 20),
        (18, "Reserva de emergência alvo", 30000),
        (21, "Parcela mensal da chácara", 1500),
        (22, "Teto mensal recomendado em dinheiro", 4000),
    )
    for row, label, value in assumptions:
        plan.cell(row, 1, label)
        plan.cell(row, 2, value)
    envelope_rows = (
        (28, "Mercado + refeições", 0),
        (29, "Casa: CPFL, BRK, internet, telefone e seguros", 650),
        (30, "Transporte", 700),
        (31, "Saúde e farmácia", 350),
        (32, "Academia", 200),
        (33, "Assinaturas", 100),
        (34, "Compras, lazer e beleza", 300),
        (35, "Pix, ajudas e outros", 300),
        (36, "Operação da chácara", 500),
        (37, "Margem para imprevistos", 900),
    )
    for row, label, cap in envelope_rows:
        plan.cell(row, 1, label)
        plan.cell(row, 3, cap)
    for offset, month in enumerate((date(2026, 9, 1), date(2026, 10, 1))):
        row = 42 + offset
        plan.cell(row, 1, month)
        plan.cell(row, 6, 1500)
        plan.cell(row, 7, 10000 if month.month == 10 else 0)

    privilege = workbook["Privilege DI"]
    privilege["A11"] = "Retorno bruto anual estimado"
    privilege["B11"] = 0.1402265
    privilege["A12"] = "IR conservador sobre os ganhos"
    privilege["B12"] = 0.225

    income = workbook["Receitas e Cenários"]
    income["K6"] = 0.06
    commissions = (
        ("COM-01", date(2026, 10, 1), 23546.16),
        ("COM-02", date(2026, 10, 1), 7111.54),
        ("COM-03", date(2027, 1, 1), 15431.25),
        ("COM-04", date(2027, 3, 1), 13558.18),
    )
    for row, (identifier, expected, gross) in enumerate(commissions, start=7):
        income.cell(row, 1, identifier)
        income.cell(row, 2, expected)
        income.cell(row, 4, gross)
    income["A16"] = date(2026, 9, 1)
    income["A17"] = date(2027, 12, 1)

    payroll = workbook["Holerites Kelly"]
    values = (
        (date(2026, 5, 1), date(2026, 6, 3), 10511.98, 5025.62, 5486.36),
        (date(2026, 6, 1), date(2026, 7, 6), 11950.73, 5414.83, 6535.90),
        (date(2026, 7, 1), date(2026, 8, 6), 11238.26, 5340.99, 5897.27),
    )
    for row, item in enumerate(values, start=6):
        for column, value in enumerate(item, start=1):
            payroll.cell(row, column, value)
        payroll.cell(row, 7, 1549.19)
        payroll.cell(row, 8, f"holerite-{row}.pdf")
    payroll["B27"] = 5026.14
    payroll["B28"] = 2296.77
    payroll["G21"] = 1675.38
    payroll["G23"] = 460.73
    payroll["G24"] = 1214.65

    card = workbook["Cartão - Dados"]
    card.append(
        [
            "Vencimento",
            "Mês da fatura",
            "Data da compra",
            "Titular/cartão",
            "Escopo",
            "Estabelecimento",
            "Valor",
            "Categoria analisada",
            "Natureza",
            "Parcelado?",
            "Parcela atual",
            "Total parcelas",
            "Categoria Itaú",
            "Arquivo-fonte",
            "Página",
        ]
    )
    card.append(
        [date(2026, 7, 8), "2026-07", date(2026, 6, 1), "KELLY (final 1234)", "Família", "Loja 01/03", 100, "Compras, casa e vestuário", "Discricionário", True, 1, 3]
    )
    card.append(
        [date(2026, 8, 8), "2026-08", date(2026, 6, 1), "KELLY (final 1234)", "Família", "Loja 02/03", 100, "Compras, casa e vestuário", "Discricionário", True, 2, 3]
    )

    nubank = workbook["Nubank - Dados"]
    nubank.append(
        ["Data", "Mês", "Descrição", "Valor", "Categoria analisada", "Natureza", "Parcelado?", "Parcela atual", "Total parcelas", "Fonte"]
    )

    bank = workbook["Banco - Dados"]
    bank.append(
        ["Data", "Mês", "Descrição", "Valor", "Categoria analisada", "Natureza", "Escopo", "Janela 6m?", "Conta no gasto familiar?", "Arquivo-fonte", "Página"]
    )
    bank.append(
        [date(2026, 8, 1), "2026-08", "CPFL", -120, "Água e energia da residência", "Essencial", "Família", True, True, "extrato.pdf", 1]
    )

    future = workbook["Parcelas Futuras"]
    future.cell(5, 6, "Mês")
    future.cell(5, 7, "Emissor")
    future.cell(5, 8, "Estabelecimento")
    future.cell(5, 10, "Total")

    review = workbook["Revisar"]
    review.append(["Título"])
    review.cell(6, 1, "Alta")
    review.cell(6, 2, "Pergunta pendente?")
    review.cell(6, 3, "Impacto no plano")
    review.cell(6, 4, "Aguardando")
    workbook.save(path)


def test_workbook_load_is_idempotent_and_installments_are_not_repeated(tmp_path) -> None:
    path = tmp_path / "plano.xlsx"
    _minimal_plan(path)
    data = parse_plan_workbook(path)
    assert len(data.transactions) == 3
    assert data.profile.investment_balance == Decimal("43875.62")
    assert data.profile.projection_end == date(2027, 12, 1)

    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        household = Household(name="Família Teste")
        db.add(household)
        db.flush()
        first = import_plan_data(db, household, data)
        db.commit()
        assert first.created["transactions"] == 3

        second = import_plan_data(db, household, data)
        db.commit()
        assert second.created["transactions"] == 0
        assert db.scalar(select(func.count(Transaction.id))) == 3
        assert _future_installments(db, household.id) == {"2026-09": Decimal("100.00")}
